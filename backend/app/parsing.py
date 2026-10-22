import io
import re
import zipfile
import unicodedata
from pathlib import Path
import pdfplumber
from docx import Document
from .catalog import SKILLS
from .security import fail

LEVEL_PATTERNS=[('advanced',r'\b(?:advanced|expert|expertise|highly proficient)\b'),('intermediate',r'\b(?:intermediate|proficient|working knowledge|comfortable|solid)\b'),('beginner',r'\b(?:beginner|basic|basics|familiar|familiarity|learning|exposure)\b')]

def read_resume(data,filename):
    ext=Path(filename).suffix.lower()
    if not data:fail('FILE_EMPTY','This file is empty.',422)
    if len(data)>5*1024*1024:fail('FILE_TOO_LARGE','Files must be 5 MB or smaller.',422)
    warnings=[]
    try:
        if ext=='.pdf':
            if not data.startswith(b'%PDF-'):fail('FILE_TYPE_UNSUPPORTED','Upload a valid PDF, DOCX or TXT file.',422)
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                if len(pdf.pages)>10:fail('TOO_MANY_PAGES','Resumes must have at most 10 pages.',422)
                text='\n'.join(p.extract_text(layout=False) or '' for p in pdf.pages)
                if len(text.strip())<50:fail('NO_TEXT_LAYER','Upload a text-based PDF or DOCX. Scanned PDFs are not supported yet.',422)
                warnings.append('Please check reading order if your PDF uses multiple columns.')
        elif ext=='.docx':
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if len(z.infolist())>2000 or sum(i.file_size for i in z.infolist())>50*1024*1024:fail('FILE_UNSAFE','The document expands beyond the supported size.',422)
                if 'word/document.xml' not in z.namelist():fail('FILE_CORRUPTED','This is not a Word document.',422)
            doc=Document(io.BytesIO(data))
            text='\n'.join([p.text for p in doc.paragraphs]+[' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows])
        elif ext=='.txt':text=data.decode('utf-8-sig')
        else:fail('FILE_TYPE_UNSUPPORTED','Upload a PDF, DOCX or UTF-8 TXT file.',422)
    except Exception as e:
        if hasattr(e,'status_code'):raise
        fail('FILE_CORRUPTED','This document could not be read. Check that it is not encrypted or corrupted.',422)
    text=unicodedata.normalize('NFKC',text).replace('\x00','')
    if len(text.strip())<80:fail('TEXT_TOO_SHORT','There is not enough readable text. Upload a complete resume.',422)
    if len(text)>60000:fail('TOO_MANY_PAGES','The document has too much text. Keep your resume within 10 pages.',422)
    return text,warnings

def extract(text):
    out=[]
    for s in SKILLS:
        mentions=[]
        for alias in sorted(s['aliases'],key=len,reverse=True):
            rx=r'(?<![\w+#.])'+re.escape(alias)+r'(?![\w+#]|\.[A-Za-z])'
            for m in re.finditer(rx,text,re.I):
                a=text.rfind('\n',0,m.start())+1;b=text.find('\n',m.end());b=len(text) if b<0 else b
                line=text[a:b].strip()
                if s['slug']=='go' and alias=='Go' and not re.search(r'golang|language|skills|programming',line,re.I):continue
                mentions.append((m.start(),m.end(),line))
        if not mentions:continue
        levels=[];evidence=[]
        for start,end,line in sorted(mentions):
            # Scope level phrases to the same semicolon/bullet-separated clause.
            clause=next((x for x in re.split(r';|\||•',line) if re.search(re.escape(text[start:end]),x,re.I)),line)
            hits=[(lvl,re.search(rx,clause,re.I)) for lvl,rx in LEVEL_PATTERNS]
            hit=next(((lvl,m.group()) for lvl,m in hits if m),None)
            if hit:levels.append(hit[0])
            rating=re.search(re.escape(text[start:end])+r'\s*(\d(?:\.\d)?)/(5|10)',clause,re.I)
            if rating and not hit:
                v=float(rating[1])/int(rating[2]);levels.append('advanced' if v>=.8 else 'intermediate' if v>=.5 else 'beginner')
            if line not in evidence:evidence.append(line[:280])
        out.append({'skill':s['slug'],'level':levels[0] if levels else None,'evidence':evidence[:3],'confidence':.9 if any(re.search(r'skills|proficient|advanced|basic',x,re.I) for x in evidence) else .78,'level_conflict':len(set(levels))>1})
    return out
